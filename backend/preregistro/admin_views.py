import csv
import re
import uuid
import unicodedata
from datetime import timedelta
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models import Count, Avg, Q
from django.db.models.functions import TruncMonth
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.http import HttpResponse
from django.http import FileResponse
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from rest_framework import permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import (
    Aspirante, Materia, PeriodoInscripcion, Inscripcion,
    DocumentoAspirante, AuditoriaSistema, ConfiguracionSistema,
)
from .configuracion_sistema import DEFAULTS as SYSTEM_CONFIG_DEFAULTS, parse_password_policy, password_meets_policy


ROLE_ALIASES = {
    'admin': 'admin', 'administrador': 'admin',
    'coordinacion': 'coordinacion', 'coordinación': 'coordinacion',
    'coordinador': 'coordinacion', 'docente': 'docente',
    'profesor': 'docente', 'director': 'director',
    'director de tesis': 'director', 'servicios': 'servicios',
    'servicios escolares': 'servicios', 'servicios_escolares': 'servicios',
    'investigador': 'investigador', 'alumno': 'alumno',
    'aspirante': 'aspirante',
}


class AdminRolePermission(permissions.BasePermission):
    """El panel administrativo queda reservado al rol admin o superusuario."""
    message = 'Se requiere el rol Administrador para acceder a este módulo.'

    def has_permission(self, request, view):
        user = request.user
        return bool(
            user and user.is_authenticated and
            (user.is_superuser or str(getattr(user, 'rol', '')).strip().lower() == 'admin')
        )


def normalizar_rol(value):
    raw = str(value or '').strip().lower()
    normalized = ''.join(ch for ch in unicodedata.normalize('NFKD', raw) if not unicodedata.combining(ch))
    return ROLE_ALIASES.get(raw, ROLE_ALIASES.get(normalized, 'aspirante'))


def rol_es_valido(value):
    raw = str(value or '').strip().lower()
    normalized = ''.join(ch for ch in unicodedata.normalize('NFKD', raw) if not unicodedata.combining(ch))
    return raw in ROLE_ALIASES or normalized in ROLE_ALIASES


def validar_password(password):
    return password_meets_policy(str(password or ''))


def validar_correo(correo):
    if len(correo) > 254:
        return False
    try:
        validate_email(correo)
    except ValidationError:
        return False
    return True


def normalizar_estado_activo(value):
    if isinstance(value, bool):
        return value
    normalized = str(value or '').strip().lower()
    if normalized in ('activo', 'active', 'true', '1'):
        return True
    if normalized in ('inactivo', 'inactive', 'false', '0'):
        return False
    return None


def usuario_es_admin_activo(user):
    return user.is_active and (user.is_superuser or normalizar_rol(user.rol) == 'admin')


def evitar_desactivar_ultimo_admin(user, activo, rol):
    actual = Aspirante.objects.only('is_active', 'is_superuser', 'rol').get(pk=user.pk)
    if usuario_es_admin_activo(actual) and (not activo or (not actual.is_superuser and normalizar_rol(rol) != 'admin')):
        admins_activos = Aspirante.objects.filter(is_active=True).filter(
            Q(is_superuser=True) | Q(rol='admin')
        ).count()
        return admins_activos <= 1
    return False


def ip_cliente(request):
    return (request.META.get('HTTP_X_FORWARDED_FOR', '').split(',')[0].strip()
            or request.META.get('REMOTE_ADDR'))


def serializar_usuario(user):
    locked_until = user.login_locked_until
    return {
        'id': user.id, 'nombre': user.nombre, 'usuario': user.usuario,
        'correo': user.correo, 'rol': user.rol, 'programa': user.programa,
        'departamento': user.departamento, 'unidad': user.unidad,
        'matricula': user.matricula, 'is_active': user.is_active,
        'is_staff': user.is_staff, 'is_superuser': user.is_superuser,
        'created_at': user.created_at.isoformat() if user.created_at else None,
        'updated_at': user.updated_at.isoformat() if user.updated_at else None,
        'proceso_estado': user.proceso_estado,
        'login_locked_until': locked_until.isoformat() if locked_until else None,
        'is_locked': bool(locked_until and locked_until > timezone.now()),
    }


def registrar_auditoria(request, accion, modelo, objeto_id='', anteriores=None, nuevos=None):
    AuditoriaSistema.objects.create(
        actor=request.user if getattr(request.user, 'is_authenticated', False) else None,
        accion=accion, modelo=modelo, objeto_id=str(objeto_id or ''),
        datos_anteriores=anteriores or {}, datos_nuevos=nuevos or {},
        ip=ip_cliente(request),
    )


def _admin_report_rows(data):
    """Convierte el resumen administrativo en filas estables para exportación."""
    keys = (
        'usuarios_total', 'usuarios_activos', 'usuarios_inactivos', 'usuarios_activos_pct',
        'alumnos', 'docentes', 'aspirantes', 'inscripciones', 'documentos',
        'documentos_pendientes', 'promedio_general', 'registros_mes_actual',
        'registros_mes_anterior', 'usuarios_variacion_pct', 'aspirantes_aceptados',
        'auditoria_ultimos_30_dias',
    )
    return [[key.replace('_', ' ').title(), data.get(key, 0)] for key in keys]


def _admin_report_xlsx(data):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Resumen'
    sheet.append(['Indicador', 'Valor'])
    for row in _admin_report_rows(data):
        sheet.append(row)
    header_fill = PatternFill('solid', fgColor='007F73')
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions['A'].width = 32
    sheet.column_dimensions['B'].width = 18
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename='reporte-administrativo.xlsx',
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _admin_report_pdf(data):
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(letter), leftMargin=0.55 * inch,
                                 rightMargin=0.55 * inch, topMargin=0.5 * inch, bottomMargin=0.55 * inch)
    styles = getSampleStyleSheet()
    story = [Paragraph('Reporte administrativo SINAC NEXT', styles['Title']),
             Paragraph('Resumen institucional generado con información del sistema.', styles['Normal']),
             Spacer(1, 12)]
    table = Table([['Indicador', 'Valor']] + _admin_report_rows(data), colWidths=[4.7 * inch, 2.0 * inch], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#007f73')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#c5d6d4')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#eef6f5')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.append(Spacer(1, 18))
    story.append(Paragraph('Documento generado por SINAC NEXT - Cinvestav Unidad Zacatenco', styles['Normal']))
    document.build(story)
    output.seek(0)
    return FileResponse(output, as_attachment=True, filename='reporte-administrativo.pdf', content_type='application/pdf')


class AdminUsuariosView(APIView):
    permission_classes = [AdminRolePermission]

    def get(self, request):
        qs = Aspirante.objects.all().order_by('-created_at')
        search = (request.query_params.get('q') or '').strip()
        rol = (request.query_params.get('rol') or '').strip().lower()
        activo = request.query_params.get('activo')
        if search:
            qs = qs.filter(nombre__icontains=search) | qs.filter(usuario__icontains=search) | qs.filter(correo__icontains=search)
        if rol:
            qs = qs.filter(rol=normalizar_rol(rol))
        if activo in ('true', 'false'):
            qs = qs.filter(is_active=activo == 'true')
        return Response({'usuarios': [serializar_usuario(u) for u in qs], 'roles': sorted({r for r in ROLE_ALIASES.values()})})

    def post(self, request):
        data = request.data
        nombre = str(data.get('nombre') or data.get('name') or '').strip()
        usuario = str(data.get('usuario') or '').strip()
        correo = str(data.get('correo') or data.get('email') or '').strip().lower()
        password = str(data.get('password') or '').strip()
        rol_solicitado = data.get('rol') or data.get('role') or 'aspirante'
        rol = normalizar_rol(rol_solicitado)
        if not nombre or not usuario or not correo or not password:
            return Response({'error': 'Nombre, usuario, correo y contraseña son obligatorios.'}, status=status.HTTP_400_BAD_REQUEST)
        if len(nombre) > 180:
            return Response({'error': 'El nombre no puede superar 180 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
        if not 3 <= len(usuario) <= 20:
            return Response({'error': 'El nombre de usuario debe tener entre 3 y 20 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
        if not validar_correo(correo):
            return Response({'error': 'Ingresa un correo electrónico válido.'}, status=status.HTTP_400_BAD_REQUEST)
        if not rol_es_valido(rol_solicitado):
            return Response({'error': 'El rol seleccionado no es válido.'}, status=status.HTTP_400_BAD_REQUEST)
        if not validar_password(password):
            return Response({'error': 'La contraseña no cumple la política configurada.'}, status=status.HTTP_400_BAD_REQUEST)
        estado_activo = normalizar_estado_activo(data.get('is_active', data.get('status', 'Activo')))
        if estado_activo is None:
            return Response({'error': 'El estado de la cuenta no es válido.'}, status=status.HTTP_400_BAD_REQUEST)
        if Aspirante.objects.filter(usuario__iexact=usuario).exists():
            return Response({'error': 'El nombre de usuario ya está registrado.'}, status=status.HTTP_400_BAD_REQUEST)
        if Aspirante.objects.filter(correo__iexact=correo).exists():
            return Response({'error': 'El correo ya está registrado.'}, status=status.HTTP_400_BAD_REQUEST)
        user = Aspirante.objects.create_user(
            usuario=usuario, correo=correo, nombre=nombre, password=password,
            curp=f'ADM{uuid.uuid4().hex[:15].upper()}',
            telefono=str(data.get('telefono') or '0000000000'),
            unidad=str(data.get('unidad') or 'Zacatenco'),
            departamento=str(data.get('departamento') or data.get('area') or 'SINAC'),
            seccion=str(data.get('seccion') or 'Administración'),
            programa=str(data.get('programa') or data.get('area') or 'SINAC NEXT'),
            estado_actual='Ciudad de México', municipio_actual='Gustavo A. Madero', direccion_actual='Cinvestav',
            estado_permanente='Ciudad de México', municipio_permanente='Gustavo A. Madero', direccion_permanente='Cinvestav',
            nombre_familiar='No aplica', parentesco='No aplica', telefono_familiar='0000000000',
            ultimo_grado='No aplica', institucion='Cinvestav', business_key=f'ADMIN-{uuid.uuid4().hex.upper()}',
            rol=rol, is_staff=rol in ('admin', 'coordinacion'),
            is_active=estado_activo,
        )
        registrar_auditoria(request, 'alta_usuario', 'Aspirante', user.id, nuevos=serializar_usuario(user))
        return Response(serializar_usuario(user), status=status.HTTP_201_CREATED)


class AdminUsuarioDetailView(APIView):
    permission_classes = [AdminRolePermission]

    def get_object(self, pk):
        try:
            return Aspirante.objects.get(pk=pk)
        except Aspirante.DoesNotExist:
            return None

    def patch(self, request, pk):
        user = self.get_object(pk)
        if not user:
            return Response(status=status.HTTP_404_NOT_FOUND)
        antes = serializar_usuario(user)
        data = request.data
        if 'nombre' in data or 'name' in data:
            nombre = str(data.get('nombre', data.get('name'))).strip()
            if not nombre:
                return Response({'error': 'El nombre no puede estar vacío.'}, status=status.HTTP_400_BAD_REQUEST)
            if len(nombre) > 180:
                return Response({'error': 'El nombre no puede superar 180 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
            user.nombre = nombre
        if 'correo' in data or 'email' in data:
            correo = str(data.get('correo', data.get('email'))).strip().lower()
            if not validar_correo(correo):
                return Response({'error': 'Ingresa un correo electrónico válido.'}, status=status.HTTP_400_BAD_REQUEST)
            if Aspirante.objects.exclude(pk=user.pk).filter(correo__iexact=correo).exists():
                return Response({'error': 'El correo ya está registrado.'}, status=status.HTTP_400_BAD_REQUEST)
            user.correo = correo
        if 'usuario' in data:
            usuario = str(data.get('usuario')).strip()
            if not 3 <= len(usuario) <= 20:
                return Response({'error': 'El nombre de usuario debe tener entre 3 y 20 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
            if Aspirante.objects.exclude(pk=user.pk).filter(usuario__iexact=usuario).exists():
                return Response({'error': 'El usuario ya está registrado.'}, status=status.HTTP_400_BAD_REQUEST)
            user.usuario = usuario
        rol_nuevo = user.rol
        if 'rol' in data or 'role' in data:
            rol_solicitado = data.get('rol', data.get('role'))
            rol_nuevo = normalizar_rol(rol_solicitado)
            if not rol_es_valido(rol_solicitado):
                return Response({'error': 'El rol seleccionado no es válido.'}, status=status.HTTP_400_BAD_REQUEST)
            user.rol = rol_nuevo
            user.is_staff = user.rol in ('admin', 'coordinacion')
        if 'area' in data:
            area = str(data.get('area') or '').strip()
            user.programa = area
            user.departamento = area
        for field in ('unidad', 'departamento', 'seccion', 'programa', 'matricula', 'proceso_estado'):
            if field in data:
                setattr(user, field, str(data.get(field) or ''))
        activo_actual = user.is_active
        activo_nuevo = activo_actual
        if 'is_active' in data or 'status' in data:
            value = data.get('is_active') if 'is_active' in data else data.get('status')
            activo_nuevo = normalizar_estado_activo(value)
            if activo_nuevo is None:
                return Response({'error': 'El estado de la cuenta no es válido.'}, status=status.HTTP_400_BAD_REQUEST)
            user.is_active = activo_nuevo
        if evitar_desactivar_ultimo_admin(user, activo_nuevo, rol_nuevo):
            return Response({'error': 'No se puede desactivar o cambiar el rol del último administrador activo.'}, status=status.HTTP_409_CONFLICT)
        if data.get('password'):
            password = str(data['password'])
            if not validar_password(password):
                return Response({'error': 'La contraseña no cumple la política configurada.'}, status=status.HTTP_400_BAD_REQUEST)
            user.set_password(password)
            user.login_failed_attempts = 0
            user.login_locked_until = None
        elif activo_nuevo and not activo_actual:
            user.login_failed_attempts = 0
            user.login_locked_until = None
        user.save()
        despues = serializar_usuario(user)
        registrar_auditoria(request, 'actualizacion_usuario', 'Aspirante', user.id, antes, despues)
        return Response(despues)

    def delete(self, request, pk):
        user = self.get_object(pk)
        if not user:
            return Response(status=status.HTTP_404_NOT_FOUND)
        if evitar_desactivar_ultimo_admin(user, False, user.rol):
            return Response({'error': 'No se puede desactivar al último administrador activo.'}, status=status.HTTP_409_CONFLICT)
        antes = serializar_usuario(user)
        user.is_active = False
        user.save(update_fields=['is_active', 'updated_at'])
        registrar_auditoria(request, 'desactivacion_usuario', 'Aspirante', user.id, antes, serializar_usuario(user))
        return Response({'detail': 'Usuario desactivado. Puede reactivarse desde la gestión de usuarios.'})


class AdminCatalogosView(APIView):
    permission_classes = [AdminRolePermission]

    def get(self, request):
        departamentos = list(Aspirante.objects.exclude(departamento='').values_list('departamento', flat=True).distinct().order_by('departamento'))
        programas = list(Aspirante.objects.exclude(programa='').values_list('programa', flat=True).distinct().order_by('programa'))
        return Response({
            'departamentos': departamentos, 'programas': programas,
            'roles': sorted(set(ROLE_ALIASES.values())),
            'materias': list(Materia.objects.order_by('clave').values('id', 'clave', 'nombre', 'creditos', 'profesor', 'horario', 'salon', 'capacidad')),
            'periodos': list(PeriodoInscripcion.objects.order_by('-apertura').values('id', 'nombre', 'apertura', 'cierre', 'activo')),
        })

    def post(self, request):
        tipo = request.data.get('tipo')
        if tipo == 'materia':
            clave = str(request.data.get('clave') or '').strip()
            nombre = str(request.data.get('nombre') or '').strip()
            if not clave or not nombre:
                return Response({'error': 'Clave y nombre son obligatorios.'}, status=status.HTTP_400_BAD_REQUEST)
            if len(clave) > 40 or len(nombre) > 200:
                return Response({'error': 'La clave admite hasta 40 caracteres y el nombre hasta 200.'}, status=status.HTTP_400_BAD_REQUEST)
            try:
                creditos = int(request.data.get('creditos') or 4)
                capacidad = int(request.data.get('capacidad') or 42)
            except (TypeError, ValueError):
                return Response({'error': 'Créditos y capacidad deben ser números válidos.'}, status=status.HTTP_400_BAD_REQUEST)
            salon = str(request.data.get('salon') or 'salon_1').strip()
            profesor = str(request.data.get('profesor') or '').strip()
            horario = str(request.data.get('horario') or '').strip()
            if len(profesor) > 180 or len(horario) > 120:
                return Response({'error': 'El profesor admite hasta 180 caracteres y el horario hasta 120.'}, status=status.HTTP_400_BAD_REQUEST)
            if creditos not in (4, 5, 7):
                return Response({'error': 'Los créditos deben ser 4, 5 o 7.'}, status=status.HTTP_400_BAD_REQUEST)
            if not 1 <= capacidad <= 32767:
                return Response({'error': 'La capacidad debe estar entre 1 y 32767.'}, status=status.HTTP_400_BAD_REQUEST)
            if salon not in dict(Materia.SALONES):
                return Response({'error': 'Selecciona un salón válido.'}, status=status.HTTP_400_BAD_REQUEST)
            materia, _ = Materia.objects.update_or_create(clave=clave, defaults={
                'nombre': nombre, 'creditos': creditos, 'salon': salon, 'capacidad': capacidad,
                'profesor': profesor, 'horario': horario,
            })
            registrar_auditoria(request, 'actualizacion_catalogo', 'Materia', materia.id, nuevos={'clave': materia.clave, 'nombre': materia.nombre})
            return Response({'id': materia.id, 'clave': materia.clave, 'nombre': materia.nombre, 'creditos': materia.creditos, 'salon': materia.salon, 'capacidad': materia.capacidad})
        if tipo == 'periodo':
            nombre = str(request.data.get('nombre') or '').strip()
            apertura = parse_datetime(str(request.data.get('apertura') or ''))
            cierre = parse_datetime(str(request.data.get('cierre') or ''))
            if not nombre or not apertura or not cierre:
                return Response({'error': 'Nombre, apertura y cierre son obligatorios y deben ser fechas válidas.'}, status=status.HTTP_400_BAD_REQUEST)
            if len(nombre) > 40:
                return Response({'error': 'El nombre del periodo no puede superar 40 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
            if timezone.is_naive(apertura): apertura = timezone.make_aware(apertura)
            if timezone.is_naive(cierre): cierre = timezone.make_aware(cierre)
            if apertura >= cierre:
                return Response({'error': 'La apertura debe ser anterior al cierre.'}, status=status.HTTP_400_BAD_REQUEST)
            activo = request.data.get('activo', True) in (True, 'true', '1', 1, 'si', 'sí')
            traslape = PeriodoInscripcion.objects.filter(activo=True, apertura__lt=cierre, cierre__gt=apertura).exclude(nombre=nombre).exists()
            if activo and traslape:
                return Response({'error': 'El periodo se traslapa con otro periodo activo.'}, status=status.HTTP_409_CONFLICT)
            periodo, _ = PeriodoInscripcion.objects.update_or_create(nombre=nombre, defaults={'apertura': apertura, 'cierre': cierre, 'activo': activo})
            registrar_auditoria(request, 'actualizacion_periodo', 'PeriodoInscripcion', periodo.id, nuevos={'nombre': periodo.nombre, 'activo': periodo.activo})
            return Response({'id': periodo.id, 'nombre': periodo.nombre, 'apertura': periodo.apertura, 'cierre': periodo.cierre, 'activo': periodo.activo})
        return Response({'error': 'Tipo de catálogo no soportado.'}, status=status.HTTP_400_BAD_REQUEST)


def resumen_alumno(alumno):
    inscripciones = list(Inscripcion.objects.filter(aspirante=alumno).select_related('materia'))
    calificaciones = [float(i.calificacion) for i in inscripciones if i.calificacion is not None]
    aprobadas = sum(1 for i in inscripciones if i.calificacion is not None and i.calificacion >= 7)
    reprobadas = sum(1 for i in inscripciones if i.calificacion is not None and i.calificacion < 7)
    return {
        'id': alumno.id, 'nombre': alumno.nombre, 'usuario': alumno.usuario,
        'matricula': alumno.matricula, 'correo': alumno.correo,
        'programa': alumno.programa, 'departamento': alumno.departamento,
        'unidad': alumno.unidad, 'proceso_estado': alumno.proceso_estado,
        'is_active': alumno.is_active, 'promedio': round(sum(calificaciones) / len(calificaciones), 2) if calificaciones else None,
        'creditos': sum(i.materia.creditos for i in inscripciones if i.calificacion is not None and i.calificacion >= 7),
        'materias_inscritas': len(inscripciones), 'materias_aprobadas': aprobadas,
        'materias_reprobadas': reprobadas,
        'expediente': ({'id': alumno.expediente_digital.id, 'folio': alumno.expediente_digital.folio, 'estado': alumno.expediente_digital.estado}
                       if hasattr(alumno, 'expediente_digital') else None),
    }


class AdminAlumnosView(APIView):
    permission_classes = [AdminRolePermission]

    def get(self, request, pk=None):
        if pk:
            alumno = Aspirante.objects.filter(pk=pk, rol='alumno', is_staff=False).first()
            if not alumno:
                return Response({'error': 'Alumno no encontrado.'}, status=status.HTTP_404_NOT_FOUND)
            inscripciones = Inscripcion.objects.filter(aspirante=alumno).select_related('materia').order_by('materia__clave')
            detalle = resumen_alumno(alumno)
            detalle['inscripciones'] = [{
                'id': item.id, 'clave': item.materia.clave, 'materia': item.materia.nombre,
                'creditos': item.materia.creditos, 'estado': item.estado,
                'parcial_1': item.parcial_1, 'parcial_2': item.parcial_2,
                'parcial_3': item.parcial_3, 'calificacion': item.calificacion,
            } for item in inscripciones]
            return Response(detalle)

        qs = Aspirante.objects.filter(rol='alumno', is_staff=False).order_by('nombre')
        search = (request.query_params.get('q') or '').strip()
        programa = (request.query_params.get('programa') or '').strip()
        departamento = (request.query_params.get('departamento') or '').strip()
        estado = (request.query_params.get('estado') or '').strip()
        if search:
            qs = qs.filter(nombre__icontains=search) | qs.filter(usuario__icontains=search) | qs.filter(matricula__icontains=search) | qs.filter(correo__icontains=search)
        if programa:
            qs = qs.filter(programa__iexact=programa)
        if departamento:
            qs = qs.filter(departamento__iexact=departamento)
        if estado:
            qs = qs.filter(proceso_estado=estado)
        return Response({'alumnos': [resumen_alumno(alumno) for alumno in qs], 'total': qs.count()})

    def patch(self, request, pk):
        alumno = Aspirante.objects.filter(pk=pk, rol='alumno', is_staff=False).first()
        if not alumno:
            return Response({'error': 'Alumno no encontrado.'}, status=status.HTTP_404_NOT_FOUND)
        antes = resumen_alumno(alumno)
        if 'matricula' in request.data:
            matricula = str(request.data.get('matricula') or '').strip()
            if len(matricula) > 30:
                return Response({'error': 'La matrícula no puede superar 30 caracteres.'}, status=status.HTTP_400_BAD_REQUEST)
            if matricula and Aspirante.objects.exclude(pk=alumno.pk).filter(matricula__iexact=matricula).exists():
                return Response({'error': 'La matrícula ya está asignada a otra cuenta.'}, status=status.HTTP_400_BAD_REQUEST)
        for field in ('programa', 'departamento', 'unidad', 'matricula', 'proceso_estado'):
            if field in request.data:
                value = matricula if field == 'matricula' else str(request.data.get(field) or '')
                setattr(alumno, field, value)
        if 'is_active' in request.data:
            active_value = normalizar_estado_activo(request.data.get('is_active'))
            if active_value is None:
                return Response({'error': 'El estado de la cuenta no es válido.'}, status=status.HTTP_400_BAD_REQUEST)
            alumno.is_active = active_value
            if active_value:
                alumno.login_failed_attempts = 0
                alumno.login_locked_until = None
        alumno.save()
        despues = resumen_alumno(alumno)
        registrar_auditoria(request, 'actualizacion_alumno', 'Aspirante', alumno.id, antes, despues)
        return Response(despues)


class AdminReportesView(APIView):
    permission_classes = [AdminRolePermission]

    def get(self, request):
        users = Aspirante.objects.all()
        alumnos = users.filter(rol='alumno')
        docentes = users.filter(rol__in=['docente', 'director', 'investigador'])
        ahora = timezone.now()
        inicio_30_dias = ahora - timedelta(days=30)
        inicio_mes_actual = ahora.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        mes_anterior = (inicio_mes_actual - timedelta(days=1)).replace(day=1)
        registros_mes_actual = users.filter(created_at__gte=inicio_mes_actual).count()
        registros_mes_anterior = users.filter(created_at__gte=mes_anterior, created_at__lt=inicio_mes_actual).count()
        variacion = (100.0 if registros_mes_actual else 0.0) if registros_mes_anterior == 0 else round(((registros_mes_actual - registros_mes_anterior) / registros_mes_anterior) * 100, 1)
        actividad = list(users.filter(created_at__gte=ahora - timedelta(days=180)).annotate(mes=TruncMonth('created_at')).values('mes').annotate(total=Count('id')).order_by('mes'))
        actividad_mensual = [{'mes': item['mes'].strftime('%b'), 'total': item['total']} for item in actividad]
        estado_aspirantes = list(users.filter(is_staff=False).values('proceso_estado').annotate(total=Count('id')).order_by('proceso_estado'))
        activos = users.filter(is_active=True).count()
        data = {
            'usuarios_total': users.count(), 'usuarios_activos': users.filter(is_active=True).count(),
            'usuarios_inactivos': users.filter(is_active=False).count(), 'alumnos': alumnos.count(),
            'cuentas_bloqueadas': users.filter(login_locked_until__gt=ahora, is_active=True).count(),
            'docentes': docentes.count(), 'aspirantes': users.filter(rol='aspirante').count(),
            'inscripciones': Inscripcion.objects.count(), 'documentos': DocumentoAspirante.objects.count(),
            'promedio_general': float(Inscripcion.objects.filter(calificacion__isnull=False).aggregate(v=Avg('calificacion'))['v'] or 0),
            'por_programa': list(alumnos.values('programa').annotate(total=Count('id')).order_by('-total')),
            'por_departamento': list(users.values('departamento').annotate(total=Count('id')).order_by('-total')),
            'inscripciones_por_estado': list(Inscripcion.objects.values('estado').annotate(total=Count('id'))),
            'usuarios_activos_pct': round((activos / users.count()) * 100, 1) if users.exists() else 0,
            'registros_mes_actual': registros_mes_actual,
            'registros_mes_anterior': registros_mes_anterior,
            'usuarios_variacion_pct': variacion,
            'aspirantes_aceptados': users.filter(rol='aspirante', proceso_estado='aceptado', is_active=True).count(),
            'documentos_pendientes': DocumentoAspirante.objects.filter(lida_notificado=False).count(),
            'auditoria_ultimos_30_dias': AuditoriaSistema.objects.filter(created_at__gte=inicio_30_dias).count(),
            'actividad_mensual': actividad_mensual,
            'estados_aspirantes': estado_aspirantes,
        }
        formato = (request.query_params.get('format') or '').lower()
        if formato in ('xlsx', 'excel'):
            return _admin_report_xlsx(data)
        if formato == 'pdf':
            return _admin_report_pdf(data)
        if formato == 'csv':
            response = HttpResponse(content_type='text/csv; charset=utf-8')
            response['Content-Disposition'] = 'attachment; filename="reporte_sinac.csv"'
            writer = csv.writer(response)
            writer.writerow(['Indicador', 'Valor'])
            for key in ('usuarios_total', 'usuarios_activos', 'usuarios_inactivos', 'usuarios_activos_pct', 'alumnos', 'docentes', 'aspirantes', 'inscripciones', 'documentos', 'documentos_pendientes', 'promedio_general', 'registros_mes_actual', 'registros_mes_anterior', 'usuarios_variacion_pct', 'aspirantes_aceptados', 'auditoria_ultimos_30_dias'):
                writer.writerow([key, data[key]])
            return response
        return Response(data)


class AdminDocentesView(APIView):
    """Consulta administrativa de docentes, investigadores y directores."""
    permission_classes = [AdminRolePermission]

    def get(self, request):
        qs = Aspirante.objects.filter(rol__in=['docente', 'director', 'investigador'], is_staff=False).order_by('nombre')
        search = (request.query_params.get('q') or '').strip()
        rol = (request.query_params.get('rol') or '').strip()
        departamento = (request.query_params.get('departamento') or '').strip()
        if search:
            qs = qs.filter(nombre__icontains=search) | qs.filter(usuario__icontains=search) | qs.filter(correo__icontains=search)
        if rol:
            qs = qs.filter(rol=rol)
        if departamento:
            qs = qs.filter(departamento__iexact=departamento)
        docentes = []
        for persona in qs:
            materias = list(Materia.objects.filter(profesor__iexact=persona.nombre).values('id', 'clave', 'nombre', 'horario', 'creditos'))
            tesistas = list(Aspirante.objects.filter(Q(director_tesis=persona) | Q(tutor_propuesto__iexact=persona.nombre)).values('id', 'nombre', 'matricula', 'programa', 'proceso_estado').distinct()) if persona.rol == 'director' else []
            docentes.append({
                'id': persona.id, 'nombre': persona.nombre, 'usuario': persona.usuario,
                'correo': persona.correo, 'rol': persona.rol, 'unidad': persona.unidad,
                'departamento': persona.departamento, 'programa': persona.programa,
                'is_active': persona.is_active, 'materias': materias,
                'total_materias': len(materias), 'tesistas': tesistas,
                'total_tesistas': len(tesistas),
                'avance_tesis_disponible': False,
            })
        return Response({'docentes': docentes, 'total': len(docentes)})


class AdminAuditoriaView(APIView):
    permission_classes = [AdminRolePermission]

    def get(self, request):
        qs = AuditoriaSistema.objects.select_related('actor').all()[:200]
        return Response([{
            'id': item.id, 'accion': item.accion, 'modelo': item.modelo, 'objeto_id': item.objeto_id,
            'usuario': item.actor.usuario if item.actor else 'Sistema', 'ip': item.ip,
            'anteriores': item.datos_anteriores, 'nuevos': item.datos_nuevos, 'created_at': item.created_at,
        } for item in qs])


class AdminConfiguracionView(APIView):
    permission_classes = [AdminRolePermission]
    defaults = SYSTEM_CONFIG_DEFAULTS

    def get(self, request):
        values = dict(self.defaults)
        values.update({
            c.clave: c.valor for c in ConfiguracionSistema.objects.filter(clave__in=self.defaults)
        })
        return Response(values)

    def put(self, request):
        unknown = set(request.data) - set(self.defaults)
        if unknown:
            return Response({'error': 'La configuración contiene campos no permitidos.'}, status=status.HTTP_400_BAD_REQUEST)
        if not request.data:
            return Response({'error': 'Envía al menos un valor para actualizar.'}, status=status.HTTP_400_BAD_REQUEST)

        values = {}
        for key, raw_value in request.data.items():
            value = str(raw_value).strip()
            if key == 'politica_password':
                policy = parse_password_policy(value)
                if policy is None:
                    return Response({
                        key: 'Usa el formato «8 caracteres, mayúscula, minúscula y número» (mínimo 8, máximo 128).'
                    }, status=status.HTTP_400_BAD_REQUEST)
                if not any((policy['uppercase'], policy['lowercase'], policy['number'])):
                    return Response({key: 'La política debe exigir al menos un tipo de carácter.'}, status=status.HTTP_400_BAD_REQUEST)
            elif key == 'bloqueo_intentos':
                try:
                    attempts = int(value)
                except ValueError:
                    return Response({key: 'Debe ser un número entero entre 1 y 20.'}, status=status.HTTP_400_BAD_REQUEST)
                if not 1 <= attempts <= 20:
                    return Response({key: 'Debe ser un número entero entre 1 y 20.'}, status=status.HTTP_400_BAD_REQUEST)
                value = str(attempts)
            elif key == 'notificaciones_activas':
                normalized = value.casefold()
                if normalized not in ('true', 'false', '1', '0', 'yes', 'no', 'si', 'sí', 'on', 'off'):
                    return Response({key: 'Envía verdadero o falso.'}, status=status.HTTP_400_BAD_REQUEST)
                value = 'true' if normalized in ('true', '1', 'yes', 'si', 'sí', 'on') else 'false'
            values[key] = value

        updated = {}
        with transaction.atomic():
            for key, value in values.items():
                config, created = ConfiguracionSistema.objects.get_or_create(
                    clave=key,
                    defaults={'valor': value, 'actualizado_por': request.user},
                )
                previous_value = config.valor
                if not created:
                    config.valor = value
                    config.actualizado_por = request.user
                    config.save(update_fields=['valor', 'actualizado_por', 'updated_at'])
                updated[key] = config.valor
                registrar_auditoria(
                    request, 'actualizacion_configuracion', 'ConfiguracionSistema', config.id,
                    anteriores={key: previous_value} if not created else {},
                    nuevos={key: config.valor},
                )
        return Response(updated)
