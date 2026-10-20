from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt


ROOT = Path(r"C:\Users\doser\OneDrive\Escritorio\Sistema SINAC NEXT")
TEMPLATE = Path(r"C:\Users\doser\Downloads\Plantilla_Entregables_SI_Cinvestav.docx")
OUTPUT = ROOT / "SINAC_NEXT_Entregable_Casos_de_Uso_v1.0.docx"


def replace_text_everywhere(document, replacements):
    for paragraph in document.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(old, new)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in paragraph.text:
                            for run in paragraph.runs:
                                run.text = run.text.replace(old, new)


def set_cell(cell, text):
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9)


def replace_table_rows(table, rows):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell(cell, value)


def insert_paragraph_before(paragraph, text="", style=None):
    # Algunas versiones de la plantilla contienen estilos duplicados con el
    # mismo nombre; insertar sin estilo evita que python-docx seleccione el
    # estilo equivocado. El contenido conserva su jerarquía mediante formato.
    new_p = paragraph.insert_paragraph_before(text)
    if style:
        for run in new_p.runs:
            run.bold = True
            run.font.size = Pt(12 if style == "Heading 2" else 11)
    return new_p


def add_case_section(anchor, case):
    title = insert_paragraph_before(anchor, f"{case['id']} · {case['name']}", "Heading 3")
    title.runs[0].bold = True
    insert_paragraph_before(anchor, f"Actor principal: {case['actor']} | Actores secundarios: {case['secondary']}")
    insert_paragraph_before(anchor, f"Objetivo: {case['objective']}")
    insert_paragraph_before(anchor, f"Precondiciones: {case['pre']}")
    insert_paragraph_before(anchor, "Flujo principal:")
    for index, step in enumerate(case['flow'], 1):
        insert_paragraph_before(anchor, f"{index}. {step}")
    insert_paragraph_before(anchor, f"Flujos alternos y excepciones: {case['alternate']}")
    insert_paragraph_before(anchor, f"Postcondiciones: {case['post']}")
    insert_paragraph_before(anchor, f"Reglas y evidencia: {case['rules']}")


def main():
    document = Document(str(TEMPLATE))
    replacements = {
        "[PRE-SI-XXX-NN]": "CINVESTAV-SI-SINAC-ED-05",
        "[1.0]": "1.0",
        "[Nombre del proyecto]": "SINAC NEXT · Sistema Integral de Gestión Académica",
        "[AAAA-MM-DD]": "2026-10-07",
        "[Unidad académica o administrativa]": "Cinvestav · Unidad Zacatenco · Coordinación Académica",
        "[Nombre(s) y adscripción]": "Equipo de desarrollo SINAC NEXT · Coordinación Académica",
        "[Nombre oficial]": "SINAC NEXT · Sistema Integral de Gestión Académica",
        "[PRE-SI-XXX]": "CINVESTAV-SI-SINAC",
        "[Departamento / coordinación / campus]": "Coordinación Académica · Servicios Escolares · Unidad Zacatenco",
        "[Nombre y cargo]": "Responsable institucional de Coordinación Académica",
        "[Nombre]": "Equipo de desarrollo SINAC NEXT",
        "[Nuevo desarrollo / reingeniería / mantenimiento / integración]": "Nuevo desarrollo e integración institucional",
        "[PMBOK / ISO 12207 / IEEE 830 / Scrum híbrido – indicar]": "BPMN 2.0 · IEEE 830 · Scrum híbrido · controles institucionales de protección de datos",
        "[AAAA-MM  a  AAAA-MM]": "2026-01 a 2026-12",
        "[Nombre completo]": "Pendiente de validación institucional",
        "[Nombre]": "Equipo de desarrollo SINAC NEXT",
        "[Describir el cambio]": "Se documentan casos de uso, actores, reglas de negocio, integraciones y criterios de aceptación.",
        "[Describa para qué se elabora este documento, a quién está dirigido y qué decisión o actividad habilita (análisis, diseño, implantación, aceptación, auditoría).]": "Este documento formaliza los casos de uso y reglas operativas de SINAC NEXT. Está dirigido a Coordinación Académica, Servicios Escolares, Administración, docentes, directores de tesis, alumnos, aspirantes y equipo técnico. Habilita la revisión de requisitos, el diseño de procesos, la construcción, las pruebas de aceptación y la auditoría funcional.",
        "[Indique qué cubre y qué queda fuera. Relacione el entregable con la fase del ciclo de vida y con otros documentos del expediente del proyecto.]": "Cubre los procesos académicos y administrativos implementados o previstos en SINAC NEXT: autenticación por rol, preregistro, admisión, expediente digital, inscripción, calificaciones, carga académica, tesis, dirección de tesis, examen predoctoral, reportes y notificaciones. Quedan fuera la definición jurídica final del aviso de privacidad, la firma institucional certificada y la operación productiva de Camunda. Corresponde a análisis y diseño, y se relaciona con el modelo de datos, BPMN, arquitectura y plan de pruebas.",
        "[Estándar o guía técnica utilizada, p. ej. IEEE 830, ISO/IEC 12207, BPMN 2.0.]": "IEEE 830 para requisitos; BPMN 2.0 para flujos; ISO/IEC 12207 como referencia del ciclo de vida; OWASP como referencia de seguridad web.",
        "[Documentos previos del mismo proyecto: charter, ERS, arquitectura, actas.]": "Repositorio SINAC NEXT, modelos Django, endpoints REST, vistas React, procesos BPMN de Camunda y documentación del modelo de base de datos.",
        "[Síntesis de una página: problema, solución propuesta, estado del entregable, decisiones pendientes y próximos pasos.]": "SINAC NEXT centraliza la gestión de aspirantes, alumnos y personal académico que anteriormente se realizaba mediante formatos, correos y registros separados. La solución integra un backend Django/REST, una interfaz React, PostgreSQL y procesos BPMN en Camunda. El entregable documenta 12 casos de uso y sus reglas de negocio, incluyendo validación de avance curricular para iniciar tesis y asignación formal de director. El estado es funcional en ambiente de desarrollo. Quedan pendientes la validación institucional de catálogos, firmas electrónicas certificadas, migración de datos históricos y pruebas UAT con usuarios reales. El siguiente paso es validar los escenarios y convertirlos en casos de prueba.",
        "[Describa el proceso, sistema o servicio actual; dolores operativos; usuarios afectados; restricciones institucionales, normativas o tecnológicas.]": "El proceso actual requiere capturar datos de aspirantes, inscripciones, calificaciones, expedientes y tesis en herramientas separadas. Esto provoca duplicidad, falta de trazabilidad, errores de captura y poca visibilidad del avance académico. Son afectados aspirantes, alumnos, docentes, directores, Coordinación Académica, Servicios Escolares y Administración. Las restricciones principales son protección de datos personales, control por roles, continuidad de expedientes, integración con Camunda y disponibilidad de infraestructura Docker/PostgreSQL.",
        "[Redacte un objetivo medible alineado al proyecto institucional.]": "Implementar y documentar un sistema institucional que permita gestionar de forma trazable al menos el 90% de los procesos académicos del ciclo aspirante–alumno–tesis, con control de roles, historial de cambios y evidencia descargable en PDF/XLSX.",
        "[Objetivo específico 1]": "Formalizar los casos de uso, actores, reglas de negocio y criterios de aceptación del sistema.",
        "[Objetivo específico 2]": "Integrar inscripciones, calificaciones, expedientes, tesis y flujos de aprobación con permisos diferenciados.",
        "[Objetivo específico 3]": "Proporcionar trazabilidad, notificaciones internas, reportes y controles de privacidad para auditoría institucional.",
        "[Qué espera del sistema]": "",
        "[Condición verificable]": "",
        "[Métrica o umbral]": "",
        "[Descripción]": "",
        "[Alto / medio / bajo]": "",
        "[Acción y responsable]": "",
    }
    replace_text_everywhere(document, replacements)

    # Control del documento, distribución y aprobaciones.
    replace_table_rows(document.tables[1], [
        ("0.1", "2026-10-07", "Equipo SINAC NEXT", "Versión inicial del entregable de casos de uso."),
        ("1.0", "2026-10-07", "Equipo SINAC NEXT", "Se incorporan actores, flujos, reglas, criterios de aceptación e integraciones."),
        ("", "", "", "Pendiente de validación por Coordinación Académica."),
    ])
    replace_table_rows(document.tables[2], [
        ("Elaboró", "Equipo de desarrollo SINAC NEXT", "Pendiente", "2026-10-07"),
        ("Revisó", "Coordinación Académica", "Pendiente", "Pendiente"),
        ("Validó (líder técnico)", "Responsable técnico SINAC NEXT", "Pendiente", "Pendiente"),
        ("Aprobó (responsable institucional)", "Responsable institucional Cinvestav", "Pendiente", "Pendiente"),
    ])
    replace_table_rows(document.tables[3], [
        ("Nombre del proyecto", "SINAC NEXT · Sistema Integral de Gestión Académica"),
        ("Código / folio institucional", "CINVESTAV-SI-SINAC"),
        ("Unidad responsable", "Coordinación Académica · Servicios Escolares · Unidad Zacatenco"),
        ("Patrocinador / sponsor", "Responsable institucional de Coordinación Académica"),
        ("Líder de proyecto", "Equipo de desarrollo SINAC NEXT"),
        ("Tipo de iniciativa", "Nuevo desarrollo e integración institucional"),
        ("Marco de referencia", "BPMN 2.0 · IEEE 830 · Scrum híbrido · protección de datos"),
        ("Periodo de ejecución", "2026-01 a 2026-12"),
    ])
    replace_table_rows(document.tables[5], [
        ("Aspirante", "Usuario externo", "Registrar expediente y consultar admisión, entrevistas y exámenes.", "Alta"),
        ("Alumno", "Usuario académico", "Inscribirse, consultar avance, registrar tesis y recibir notificaciones.", "Alta"),
        ("Docente / Director", "Usuario académico", "Gestionar grupos, entrevistas, evaluaciones, tesis y dictámenes.", "Alta"),
        ("Coordinación Académica", "Responsable de proceso", "Validar expedientes, periodos, cargas, tesis y asignaciones.", "Alta"),
        ("Servicios Escolares", "Custodio administrativo", "Validar inscripciones, historial, constancias y estatus.", "Alta"),
        ("Administrador / TI", "Soporte y seguridad", "Gestionar usuarios, catálogos, auditoría y configuración.", "Media"),
    ])
    replace_table_rows(document.tables[6], [
        ("RF-01", "El sistema deberá autenticar y direccionar al usuario según su rol.", "Funcional", "Alta", "Cada rol accede únicamente a su panel autorizado."),
        ("RF-02", "El sistema deberá registrar y consultar aspirantes y alumnos.", "Funcional", "Alta", "Se puede crear, buscar, editar y consultar el expediente autorizado."),
        ("RF-03", "El sistema deberá controlar inscripciones por periodo activo.", "Funcional", "Alta", "El alumno solo elige materias dentro de apertura y cierre."),
        ("RF-04", "El docente deberá capturar parciales y calificación final.", "Funcional", "Alta", "La final determina aprobado, reprobado o cursando."),
        ("RF-05", "El sistema deberá validar elegibilidad e iniciar solicitud de director de tesis.", "Funcional", "Alta", "Solo se habilita al completar materias, créditos y periodos."),
        ("RF-06", "Coordinación deberá asignar formalmente al director de tesis.", "Funcional", "Alta", "La asignación actualiza alumno, solicitud, proyecto y seguimiento."),
        ("RF-07", "El sistema deberá programar y evaluar el examen predoctoral.", "Funcional", "Alta", "Director o Coordinación registra fecha, estado, calificación y retroalimentación."),
        ("RF-08", "El sistema deberá generar reportes PDF/XLSX y bitácoras.", "Funcional", "Media", "El archivo conserva filtros, fecha y usuario generador."),
        ("RNF-01", "Control de acceso por rol y objeto académico.", "No funcional", "Alta", "Las API responden 403 ante acceso no autorizado."),
        ("RNF-02", "Protección de datos personales y trazabilidad.", "No funcional", "Alta", "Cambios relevantes registran actor, fecha, modelo y valores."),
    ])
    replace_table_rows(document.tables[7], [
        ("R-01", "Catálogos de programas y materias incompletos", "Alto", "Media", "Coordinación valida configuración del plan antes de abrir inscripciones."),
        ("R-02", "Calificaciones finales no capturadas o no cerradas", "Alto", "Media", "El avance conserva estado Cursando hasta recibir calificación final."),
        ("R-03", "Motor Camunda no disponible", "Medio", "Media", "Registrar estado pendiente, reintentar y conservar evidencia local."),
        ("R-04", "Asignación de director incompatible con programa", "Alto", "Baja", "Validar rol, disponibilidad, departamento y línea de investigación."),
        ("R-05", "Tratamiento indebido de datos personales", "Alto", "Baja", "Aplicar mínimo privilegio, aviso de privacidad y bitácora."),
    ])
    replace_table_rows(document.tables[8], [
        ("SINAC NEXT", "Sistema Integral de Gestión Académica."),
        ("LIDA", "Línea de integración y desarrollo académico asociada a los procesos de admisión y tesis."),
        ("ERS", "Especificación de requisitos de software."),
        ("RF / RNF", "Requisito funcional / no funcional."),
        ("UAT", "Pruebas de aceptación de usuario."),
        ("BPMN", "Notación para modelar procesos de negocio; Camunda ejecuta los flujos publicados."),
        ("Retícula", "Plan de materias, créditos y periodos que integran un programa académico."),
        ("Expediente digital", "Conjunto controlado de documentos, firmas, estados y evidencias del alumno."),
    ])

    # Insert detailed use cases immediately before architecture section.
    anchor = next(p for p in document.paragraphs if p.text.strip() == "4.6 Arquitectura y diseño (si aplica)")
    insert_paragraph_before(anchor, "4.5.1 Catálogo detallado de casos de uso", "Heading 2")
    insert_paragraph_before(anchor, "Los casos siguientes describen el comportamiento esperado de SINAC NEXT. Los estados y permisos indicados son reglas de negocio y deben comprobarse en pruebas UAT.")
    cases = [
        {"id":"CU-01","name":"Autenticar usuario y direccionar por rol","actor":"Aspirante, alumno, docente, director, coordinación, servicios escolares o administrador","secondary":"API de autenticación y bitácora","objective":"Validar credenciales, emitir sesión y presentar únicamente el panel autorizado.","pre":"Cuenta activa, credenciales válidas y rol asignado.","flow":["El usuario captura usuario/correo y contraseña.","El backend valida credenciales y estado de la cuenta.","Se emite token de sesión y se identifica el rol institucional.","La interfaz direcciona al panel correspondiente.","Cada consulta posterior valida el token y el alcance del rol."],"alternate":"Credenciales inválidas muestran error; cuenta inactiva o rol no reconocido bloquea el acceso y se registra el evento.","post":"Sesión iniciada o intento rechazado con evidencia de auditoría.","rules":"No se revela si falló usuario o contraseña; las API deben devolver 401/403 correctamente."},
        {"id":"CU-02","name":"Registrar expediente de aspirante","actor":"Aspirante","secondary":"Coordinación Académica, Camunda, notificaciones","objective":"Crear un preregistro completo y trazable para admisión.","pre":"El aviso de privacidad fue aceptado; CURP, correo, teléfono y contraseña cumplen validaciones.","flow":["El aspirante captura datos personales, domicilio, escolaridad y programa.","El sistema valida CURP, correo, teléfono, contraseña y duplicados.","Se crea el usuario con estado pendiente y clave de negocio.","Se inicia el flujo de preregistro en Camunda cuando está disponible.","El aspirante consulta documentos, seguimiento, entrevista y examen."],"alternate":"Datos duplicados se rechazan; archivo inválido queda pendiente de corrección; Camunda no disponible conserva el expediente local y registra incidencia.","post":"Expediente creado, seguimiento inicial y acceso al portal de aspirante.","rules":"Los datos personales se muestran solo a roles autorizados y se conserva el aviso de privacidad."},
        {"id":"CU-03","name":"Evaluar y resolver admisión","actor":"Coordinación Académica / Colegio de Profesores","secondary":"Aspirante, docente evaluador, Camunda","objective":"Revisar expediente, examen, entrevista y dictamen de admisión.","pre":"Existe expediente de aspirante y documentación mínima disponible.","flow":["Coordinación filtra solicitudes por estado.","Consulta expediente, documentos y seguimiento.","Programa o revisa examen y entrevista.","El Colegio registra dictamen favorable, no favorable o en revisión.","Coordinación acepta o rechaza y registra observaciones.","El sistema notifica el cambio y actualiza el proceso."],"alternate":"Si faltan documentos, se mantiene en revisión; si se rechaza, el aspirante consulta el motivo y no se convierte a alumno.","post":"Estado de admisión actualizado con dictamen, fecha y actor.","rules":"Las acciones de aceptar/rechazar solo aparecen para usuarios aspirantes."},
        {"id":"CU-04","name":"Aceptar aspirante y crear expediente digital","actor":"Coordinación Académica","secondary":"Servicios Escolares, aspirante aceptado, expediente digital","objective":"Convertir al aspirante aceptado en alumno y generar su expediente institucional.","pre":"Dictamen favorable y expediente aceptado.","flow":["Coordinación confirma aceptación.","El sistema crea o conserva la cuenta sin duplicarla.","Genera matrícula o número de control.","Crea expediente digital y registra documentos iniciales.","Asocia programa, grado, unidad y director pendiente.","Envía notificación interna al nuevo alumno."],"alternate":"Si matrícula, CURP o clave de negocio ya existen, se reutiliza la cuenta o se solicita corrección; nunca se crean duplicados.","post":"Alumno activo, expediente digital creado y seguimiento de conversión registrado.","rules":"La conversión es idempotente y debe quedar auditada."},
        {"id":"CU-05","name":"Inscribir materias dentro del periodo","actor":"Alumno","secondary":"Coordinación Académica, Servicios Escolares","objective":"Permitir selección y validación de materias únicamente durante un periodo abierto.","pre":"Alumno activo; periodo de inscripción abierto; catálogo y plan disponibles.","flow":["Coordinación publica apertura y cierre del periodo.","El alumno consulta materias, créditos, docente y horario.","Selecciona una materia.","El sistema valida que no exista inscripción duplicada y registra la solicitud.","Se muestra estado cursando y se genera boleta/horario descargable."],"alternate":"Fuera de fechas se bloquea la acción; materia duplicada o sin cupo produce mensaje; cambios quedan sujetos a validación administrativa.","post":"Inscripción relacionada con alumno, materia y periodo.","rules":"La fecha se valida en servidor, no solo en la interfaz."},
        {"id":"CU-06","name":"Capturar parciales y calificación final","actor":"Docente","secondary":"Alumno, Coordinación Académica","objective":"Registrar evaluación por parciales y cerrar la materia con una calificación final.","pre":"Docente asociado a la materia y alumno inscrito.","flow":["Docente consulta sus materias y alumnos por grupo.","Captura o actualiza parcial 1, 2 y 3.","El sistema valida rango 0–10 y conserva cambios.","Al cerrar, captura calificación final.","El sistema determina aprobada (>=7), reprobada (<7) o cursando si falta final.","El alumno visualiza calificaciones y alertas de riesgo."],"alternate":"Un docente no puede modificar materias ajenas; una final faltante mantiene Cursando; correcciones deben quedar auditadas.","post":"Calificación registrada, estado académico calculado y avance actualizado.","rules":"La final es la única que determina aprobación definitiva."},
        {"id":"CU-07","name":"Aprobar carga académica multietapa","actor":"Coordinación Académica","secondary":"Camunda, docentes, Servicios Escolares","objective":"Crear, revisar, aprobar y publicar la carga académica de un periodo.","pre":"Catálogo de materias, docentes y horarios disponible; proceso BPMN desplegado.","flow":["Coordinación crea carga con periodo.","El backend inicia instancia Camunda y guarda su identificador.","Se envía a revisión y se consulta tarea activa.","Coordinación aprueba o rechaza con comentario.","Una carga aprobada se publica para consulta e inscripción.","El sistema muestra estado y conexión Camunda."],"alternate":"Si Camunda no responde se conserva la carga local como pendiente; no se publica hasta completar la aprobación.","post":"Carga publicada o rechazada con trazabilidad.","rules":"Estados: borrador, en revisión, aprobada, publicada o rechazada."},
        {"id":"CU-08","name":"Validar expediente y firmar electrónicamente","actor":"Director de Tesis / Servicios Escolares","secondary":"Alumno, expediente digital, bitácora","objective":"Revisar documentos y registrar firma del responsable autorizado.","pre":"Expediente digital generado y usuario con rol firmante.","flow":["El responsable consulta expedientes de su alcance.","Revisa documentos y estado de firmas.","Confirma firma electrónica con huella y rol.","El sistema evita firma duplicada y registra fecha.","El alumno consulta el estado de su expediente."],"alternate":"Documento faltante mantiene expediente pendiente; firma repetida se rechaza sin duplicar registros.","post":"Firma asociada al expediente, usuario, rol y marca de tiempo.","rules":"La firma implementada es de control interno; la certificación jurídica debe definirse institucionalmente."},
        {"id":"CU-09","name":"Validar elegibilidad y solicitar director de tesis","actor":"Alumno","secondary":"Coordinación Académica, catálogo de programa","objective":"Habilitar el inicio de tesis solo al completar requisitos del grado.","pre":"Programa configurado con grado, periodicidad, periodos, créditos y materias obligatorias.","flow":["El sistema identifica maestría (2 años) o doctorado (4 años).","Calcula semestres o cuatrimestres concluidos.","Suma materias obligatorias aprobadas y créditos.","Muestra requisitos pendientes o habilita la solicitud.","El alumno captura línea de investigación y justificación.","Se crea solicitud pendiente para Coordinación."],"alternate":"Si faltan materias, créditos o periodos, el formulario de solicitud no se habilita y se explica cada pendiente.","post":"Solicitud de dirección de tesis registrada o requisitos mostrados.","rules":"Maestría: 4 semestres/6 cuatrimestres; doctorado: 8 semestres/12 cuatrimestres. Aprobación mínima: 7."},
        {"id":"CU-10","name":"Asignar Director de Tesis","actor":"Coordinación Académica","secondary":"Alumno, Director de Tesis, ProyectoTesis","objective":"Asignar formalmente un director compatible y crear el vínculo académico.","pre":"Solicitud de alumno aprobable y director activo con rol válido.","flow":["Coordinación consulta solicitudes pendientes y directores disponibles.","Revisa programa, departamento, línea de investigación y carga de tesistas.","Selecciona director y registra observaciones.","El sistema actualiza alumno.director_tesis.","Se registra seguimiento para alumno y director.","El proyecto de tesis queda disponible para registro y revisión."],"alternate":"Director inactivo o rol incompatible se rechaza; una solicitud rechazada conserva motivo; no se asignan cuentas duplicadas.","post":"Relación alumno-director formal, auditable y visible para ambos paneles.","rules":"El director solo puede consultar tesistas asignados; Coordinación mantiene autoridad de asignación."},
        {"id":"CU-11","name":"Registrar proyecto y examen predoctoral","actor":"Alumno / Director de Tesis","secondary":"Coordinación Académica","objective":"Gestionar proyecto de tesis y, para doctorado, programar y resolver examen predoctoral.","pre":"Alumno elegible y director asignado; proyecto en estado borrador o revisión.","flow":["El alumno registra título, resumen, objetivos, metodología y línea.","Envía proyecto a revisión.","Director y Coordinación dictaminan aprobado, observado o rechazado.","Para doctorado, el director programa examen predoctoral.","El director registra estado, calificación y retroalimentación.","El alumno consulta resultado y avance."],"alternate":"Sin director asignado no se registra proyecto; el alumno no puede modificar el resultado del examen; el examen predoctoral no usa el flujo automático de admisión.","post":"Proyecto dictaminado y examen predoctoral registrado cuando aplica.","rules":"Maestría no requiere examen predoctoral; doctorado sí lo contempla antes de la defensa final."},
        {"id":"CU-12","name":"Consultar expediente, reportes y notificaciones","actor":"Servicios Escolares / Administración","secondary":"Todos los usuarios autorizados","objective":"Consultar información institucional y generar evidencia sin alterar responsabilidades académicas.","pre":"Usuario autenticado y permisos vigentes.","flow":["El usuario filtra alumnos, aspirantes, docentes o directores.","Consulta expediente, historial, estatus, créditos o reportes.","Genera PDF/XLSX con fecha y usuario.","Consulta notificaciones internas y bitácora.","El sistema registra la consulta o modificación relevante."],"alternate":"Acceso fuera de alcance responde 403; datos incompletos se muestran como pendientes, no se inventan.","post":"Consulta o archivo generado con trazabilidad.","rules":"Servicios Escolares consulta y valida calificaciones; no modifica directamente una calificación docente sin procedimiento autorizado."},
    ]
    for case in cases:
        add_case_section(anchor, case)

    # Replace selected section paragraphs with project-specific content.
    paragraph_replacements = {
        "Datos personales o sensibles involucrados: [sí / no; catálogo].": "Datos personales o sensibles involucrados: sí. CURP, nombre, matrícula, correo, teléfono, domicilio, trayectoria académica, documentos, firmas y datos de acceso.",
        "Base legal y aviso de privacidad: [referencia].": "Base legal y aviso de privacidad: aviso de privacidad de Cinvestav disponible desde el portal; su versión jurídica definitiva debe ser validada por la Unidad de Transparencia.",
        "Controles de acceso, bitácoras y respaldo: [describir].": "Controles: autenticación con token, permisos por rol y objeto, bitácora de auditoría, restricciones de API, respaldos PostgreSQL y control de archivos.",
        "Retención y disposición de la información: [política].": "Retención: conservar expedientes académicos conforme a la política institucional y disposiciones archivísticas; definir plazos con Servicios Escolares y Transparencia.",
        "[Estrategia de pruebas (unitarias, integración, UAT), responsables, ambientes, criterios de salida y evidencia a resguardar en el expediente.]": "Estrategia: pruebas unitarias de validadores y servicios; integración de API, PostgreSQL y Camunda; UAT por rol con Coordinación, docentes, alumnos y Servicios Escolares. Criterio de salida: escenarios críticos exitosos, sin errores de autorización y evidencia de respuestas, capturas, reportes y bitácoras.",
        "[Ventana de salida a producción, migración de datos, rollback, usuarios piloto, materiales de capacitación y soporte de primer nivel.]": "Implantación: piloto con Coordinación Académica y Servicios Escolares; migración controlada de catálogos y usuarios; respaldo previo y rollback por imagen/base de datos; capacitación por rol; soporte inicial mediante bitácora de incidencias.",
        "[Estado del entregable, acuerdos, pendientes y fecha estimada del siguiente hito.]": "Estado: entregable en versión 1.0 para revisión. Acuerdos: validar catálogos, roles, reglas de grado y firma institucional. Pendientes: UAT, datos históricos, política de retención y certificación de firma. Siguiente hito: revisión institucional y matriz de pruebas.",
    }
    replace_text_everywhere(document, paragraph_replacements)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Nomenclatura sugerida:"):
            paragraph.text = "Nomenclatura aplicada: CINVESTAV-SI-SINAC-ED-05-v1.0. El código ED-05 identifica el modelo de procesos y casos de uso."
        elif paragraph.text.startswith("[Describa componentes"):
            paragraph.text = "Arquitectura: frontend React con paneles por rol; backend Django REST Framework; PostgreSQL como repositorio; JWT para sesión; Camunda REST para preregistro y carga académica; generación PDF/XLSX; almacenamiento controlado de documentos; ambientes de desarrollo, pruebas y producción separados. Las decisiones priorizan mínimo privilegio, trazabilidad, idempotencia y validación en servidor. Los BPMN existentes se incluyen como evidencia técnica del expediente."
        elif paragraph.text.startswith("Anexo D."):
            paragraph.text = "Anexo D. Matriz de trazabilidad entre casos de uso, endpoints REST, vistas por rol y procesos BPMN."
    document.core_properties.title = "SINAC NEXT · Casos de uso y procesos institucionales"
    document.core_properties.subject = "Entregable ED-05 · Modelo de procesos y casos de uso"
    document.core_properties.author = "Equipo de desarrollo SINAC NEXT"
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == '__main__':
    main()
